import argparse
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


W, H = 13.333, 7.5
CREAM = RGBColor(250, 247, 240)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(43, 42, 38)
MUTED = RGBColor(100, 96, 88)
BROWN = RGBColor(102, 67, 48)
ROSE = RGBColor(192, 78, 92)
GREEN = RGBColor(46, 105, 88)
LINE = RGBColor(220, 214, 204)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_table(slide, data, x, y, w, h, col_widths, header_fill):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(
        rows, cols, Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for index, width in enumerate(col_widths):
        table.columns[index].width = Inches(width)
    for row in range(rows):
        for col in range(cols):
            cell = table.cell(row, col)
            cell.text = data[row][col]
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                header_fill if row == 0
                else WHITE if row % 2 else RGBColor(244, 241, 235)
            )
            cell.text_frame.word_wrap = True
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(10.5 if row else 10)
                    run.font.bold = row == 0 or col == 0
                    run.font.color.rgb = WHITE if row == 0 else INK


def add_annex_header(slide, number, section, title, subtitle):
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = CREAM
    add_text(slide, f"{number:02d}  {section}", 0.62, 0.28, 4.4, 0.30,
             10, ROSE, True)
    add_text(slide, title, 0.62, 0.72, 12.0, 0.55, 25, INK, True)
    add_text(slide, subtitle, 0.64, 1.33, 12.0, 0.40, 13, MUTED)
    marker = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.75), Inches(0.76), Inches(0.05)
    )
    marker.fill.solid()
    marker.fill.fore_color.rgb = ROSE
    marker.line.fill.background()
    add_text(slide, "Diseño de Software · Semana 3", 0.62, 7.14, 4.0, 0.18,
             9, MUTED)
    add_text(slide, str(number), 12.25, 7.10, 0.45, 0.20,
             9, MUTED, align=PP_ALIGN.RIGHT)


def build(pdf_path: Path, output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "Proceso Unificado aplicado a Dolce‑Vita"
    prs.core_properties.subject = "Presentación oficial con tablas de referencia"
    prs.core_properties.author = "Cristopher Alexander Bañares Suarez"

    with tempfile.TemporaryDirectory(prefix="dolce-vita-pdf-") as temp:
        image_prefix = Path(temp) / "pagina"
        subprocess.run(
            [
                "pdftoppm", "-jpeg", "-jpegopt", "quality=92", "-r", "150",
                str(pdf_path), str(image_prefix)
            ],
            check=True,
        )
        pages = sorted(Path(temp).glob("pagina-*.jpg"))
        if len(pages) != 16:
            raise RuntimeError(f"Se esperaban 16 páginas y se encontraron {len(pages)}")
        for page in pages:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(
                str(page), Inches(0), Inches(0), width=Inches(W), height=Inches(H)
            )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_annex_header(
        slide, 17, "ANEXO · ROLES", "Tabla de roles aplicada a Dolce‑Vita",
        "Responsabilidades definidas según el material oficial de la asignatura."
    )
    role_data = [
        ["Rol", "Responsabilidad principal", "Aplicación en Dolce‑Vita"],
        ["Analista", "Captura y documenta requerimientos.",
         "Define reglas de registro, compra, stock, alertas y boleta."],
        ["Diseñador/a", "Crea la arquitectura y los diagramas UML.",
         "Modela usuarios, productos, carro, inventario y ventas."],
        ["Programador/a", "Implementa el código del sistema.",
         "Programa módulos de acceso, compra, stock y boleta."],
        ["Tester", "Verifica que el software cumpla los requisitos.",
         "Prueba compra, inventario, alertas y emisión de boleta."],
        ["Administrador/a de proyecto", "Planifica, coordina y supervisa.",
         "Organiza iteraciones, recursos, plazos, riesgos e hitos."],
    ]
    add_table(slide, role_data, 0.62, 2.04, 12.08, 4.45,
              [2.32, 4.15, 5.61], BROWN)
    add_text(
        slide,
        "Fuente: Santo Tomás Online. Tabla de roles y responsabilidades en el Proceso Unificado.",
        0.75, 6.63, 11.75, 0.22, 9, MUTED, align=PP_ALIGN.CENTER
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_annex_header(
        slide, 18, "ANEXO · TABLA DE FASES", "Fases del Proceso Unificado",
        "Actividades y artefactos vinculados al caso de ventas en línea."
    )
    phase_data = [
        ["Fase", "Actividades principales", "Artefactos", "Dolce‑Vita"],
        ["Inicio", "Requisitos, alcance, planificación y costos.",
         "Casos de uso y visión.", "Priorizar acceso, compra, stock y boleta."],
        ["Elaboración", "Análisis, diseño, arquitectura, iteraciones y riesgos.",
         "UML y prototipos.", "Validar arquitectura de ventas y stock."],
        ["Construcción", "Implementación, integración, pruebas y revisiones.",
         "Código y manual técnico.", "Construir módulos por incremento."],
        ["Transición", "Pruebas finales, despliegue, capacitación y estabilización.",
         "Plan y manual de usuario.", "Piloto, correcciones y puesta en marcha."],
    ]
    add_table(slide, phase_data, 0.43, 2.02, 12.47, 4.45,
              [1.42, 4.10, 2.84, 4.11], GREEN)
    add_text(
        slide,
        "Fuente: Santo Tomás Online. Tabla de fases del Proceso Unificado de Desarrollo de Software.",
        0.75, 6.63, 11.75, 0.22, 9, MUTED, align=PP_ALIGN.CENTER
    )

    prs.save(output_path)
    print(f"Presentación creada desde el PDF oficial: {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("pdf", type=Path, help="PDF oficial de 16 páginas")
    parser.add_argument(
        "-o", "--output", type=Path,
        default=Path("Presentacion_Dolce_Vita_Oficial_con_Tablas.pptx")
    )
    args = parser.parse_args()
    build(args.pdf.resolve(), args.output.resolve())
