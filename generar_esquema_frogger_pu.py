from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


PPTX_OUTPUT = Path("S2_Bañares_Proceso_Unificado_Frogger.pptx")
PDF_OUTPUT = Path("S2_Bañares_Proceso_Unificado_Frogger.pdf")

NAVY = "12263A"
DARK = "243B53"
MID = "627D98"
LIGHT = "F4F7FA"
WHITE = "FFFFFF"
GREEN = "39A845"
ORANGE = "F59E42"
BLUE = "2F80ED"
PURPLE = "8B5CF6"
PALE_GREEN = "EAF7EB"
PALE_ORANGE = "FFF3E7"
PALE_BLUE = "EAF3FF"
PALE_PURPLE = "F2ECFF"
LINE = "D9E2EC"

phases = [
    {
        "number": "01",
        "name": "INICIO",
        "color": GREEN,
        "pale": PALE_GREEN,
        "objective": "Definir alcance y comprobar la viabilidad del rediseño.",
        "activities": [
            "Acordar estética retro, fecha límite y versiones PC/móvil.",
            "Identificar casos de uso: iniciar partida, mover la rana y cruzar el nivel.",
            "Detectar riesgos de controles táctiles y compatibilidad.",
        ],
        "artifacts": ["Visión", "Casos de uso iniciales", "Lista de riesgos"],
    },
    {
        "number": "02",
        "name": "ELABORACIÓN",
        "color": ORANGE,
        "pale": PALE_ORANGE,
        "objective": "Validar la arquitectura y reducir los riesgos principales.",
        "activities": [
            "Refinar movimiento, colisiones, vidas, puntaje y pausa.",
            "Prototipar controles de teclado y pantalla táctil.",
            "Diseñar motor, niveles, obstáculos e interfaz multiplataforma.",
        ],
        "artifacts": ["Diagrama de clases", "Secuencia", "Prototipo jugable"],
    },
    {
        "number": "03",
        "name": "CONSTRUCCIÓN",
        "color": BLUE,
        "pale": PALE_BLUE,
        "objective": "Desarrollar el juego mediante incrementos funcionales.",
        "activities": [
            "I1: movimiento, mapa y colisiones básicas.",
            "I2: tráfico, río, vidas, puntaje y dificultad.",
            "I3: interfaz retro, audio y controles PC/móvil; integrar y probar.",
        ],
        "artifacts": ["Código fuente", "Builds jugables", "Pruebas unitarias"],
    },
    {
        "number": "04",
        "name": "TRANSICIÓN",
        "color": PURPLE,
        "pale": PALE_PURPLE,
        "objective": "Entregar una versión estable y aceptada por usuarios.",
        "activities": [
            "Probar compatibilidad, rendimiento y controles en PC/móvil.",
            "Realizar beta con jugadores y corregir observaciones.",
            "Publicar paquetes finales y entregar soporte básico.",
        ],
        "artifacts": ["Versión final", "Informe de aceptación", "Manual de usuario"],
    },
]


def rgb(value):
    return RGBColor.from_string(value)


def add_shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else kind,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def text_in(
    shape,
    text,
    size,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.05,
):
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, **kwargs):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    text_in(shape, text, **kwargs)
    return shape


def add_multiline(slide, x, y, w, h, lines, margin=0.14, fill=WHITE, line=LINE):
    shape = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line, True)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for index, item in enumerate(lines):
        text, size, color, bold, after = item
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, LIGHT)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.12, GREEN)
    add_text(
        slide,
        0.48,
        0.28,
        9.4,
        0.48,
        "Proceso Unificado aplicado al rediseño de Frogger",
        size=23,
        color=NAVY,
        bold=True,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        0.5,
        0.78,
        8.8,
        0.26,
        "Flujo iterativo para modernizar el juego retro y llevarlo a PC y dispositivos móviles",
        size=10,
        color=MID,
    )
    name_box = add_shape(slide, MSO_SHAPE.RECTANGLE, 10.1, 0.32, 2.72, 0.58, NAVY, NAVY, True)
    text_in(
        name_box,
        "CRISTOPHER ALEXANDER\nBAÑARES SUAREZ",
        size=8.8,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0.08,
    )

    card_x = [0.42, 3.65, 6.88, 10.11]
    for idx, (phase, x) in enumerate(zip(phases, card_x)):
        # Card
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.47, 2.80, 4.62, WHITE, LINE, True)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 1.47, 2.80, 0.72, phase["color"], phase["color"], True)
        add_text(
            slide,
            x + 0.15,
            1.58,
            0.45,
            0.36,
            phase["number"],
            size=12,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            x + 0.7,
            1.57,
            1.8,
            0.38,
            phase["name"],
            size=13,
            color=WHITE,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            x + 0.17,
            2.35,
            2.45,
            0.56,
            phase["objective"],
            size=9.5,
            color=NAVY,
            bold=True,
        )
        add_text(slide, x + 0.17, 2.97, 1.4, 0.22, "ACTIVIDADES", size=7.5, color=phase["color"], bold=True)
        bullets = "\n".join(f"• {activity}" for activity in phase["activities"])
        bullet_box = add_text(
            slide,
            x + 0.15,
            3.22,
            2.47,
            1.83,
            bullets,
            size=8.6,
            color=DARK,
            valign=MSO_ANCHOR.TOP,
            margin=0.02,
        )
        for p in bullet_box.text_frame.paragraphs:
            p.space_after = Pt(5)
        add_text(slide, x + 0.17, 5.12, 1.4, 0.22, "ARTEFACTOS", size=7.5, color=phase["color"], bold=True)
        for tag_idx, artifact in enumerate(phase["artifacts"]):
            pill = add_shape(
                slide,
                MSO_SHAPE.RECTANGLE,
                x + 0.15,
                5.38 + tag_idx * 0.22,
                2.46,
                0.19,
                phase["pale"],
                phase["pale"],
                True,
            )
            text_in(
                pill,
                artifact,
                size=6.8,
                color=phase["color"],
                bold=True,
                valign=MSO_ANCHOR.MIDDLE,
                margin=0.015,
            )

        # Sequential dependency arrows
        if idx < 3:
            arrow = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + 2.82),
                Inches(1.82),
                Inches(card_x[idx + 1] - 0.02),
                Inches(1.82),
            )
            arrow.line.color.rgb = rgb(MID)
            arrow.line.width = Pt(2)
            arrow.line.end_arrowhead = True

    labels = [
        ("Alcance aprobado", 3.07),
        ("Arquitectura base", 6.29),
        ("Incrementos listos", 9.48),
    ]
    for label, x in labels:
        add_text(slide, x, 1.14, 1.34, 0.21, label, size=6.5, color=MID, bold=True, align=PP_ALIGN.CENTER)

    # Iteration band
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.42, 6.28, 12.49, 0.84, NAVY, NAVY, True)
    add_text(slide, 0.67, 6.47, 1.25, 0.25, "ITERACIONES", size=9, color=WHITE, bold=True)
    add_text(
        slide,
        1.78,
        6.42,
        5.65,
        0.34,
        "Requisitos  →  diseño  →  implementación  →  pruebas  →  incremento jugable",
        size=9.4,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        7.65,
        6.40,
        4.92,
        0.38,
        "↶  La beta y las pruebas PC/móvil devuelven ajustes a Construcción y Elaboración",
        size=8.2,
        color="CFE8D1",
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        0.45,
        7.23,
        12.43,
        0.14,
        "Diseño de Software · Semana 2  |  El PU valida temprano y permite adaptar requisitos; no espera al final como un proceso en cascada.",
        size=6.7,
        color=MID,
        align=PP_ALIGN.CENTER,
    )
    prs.save(PPTX_OUTPUT)


pdfmetrics.registerFont(TTFont("Arimo", "/usr/share/fonts/truetype/croscore/Arimo-Regular.ttf"))
pdfmetrics.registerFont(TTFont("Arimo-Bold", "/usr/share/fonts/truetype/croscore/Arimo-Bold.ttf"))


def draw_wrapped(c, text, x, y, max_width, font_name, font_size, color, leading, max_lines=None):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if pdfmetrics.stringWidth(candidate, font_name, font_size) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    if max_lines:
        lines = lines[:max_lines]
    c.setFont(font_name, font_size)
    c.setFillColor(HexColor(f"#{color}"))
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def rounded(c, x, y, w, h, fill, stroke=None, radius=8):
    c.setFillColor(HexColor(f"#{fill}"))
    c.setStrokeColor(HexColor(f"#{stroke or fill}"))
    c.roundRect(x, y, w, h, radius, fill=1, stroke=1)


def build_pdf():
    width, height = landscape((13.333 * inch, 7.5 * inch))
    c = canvas.Canvas(str(PDF_OUTPUT), pagesize=(width, height))
    c.setTitle("Proceso Unificado aplicado al rediseño de Frogger")
    c.setAuthor("Cristopher Alexander Bañares Suarez")
    c.setFillColor(HexColor(f"#{LIGHT}"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(HexColor(f"#{GREEN}"))
    c.rect(0, height - 0.12 * inch, width, 0.12 * inch, fill=1, stroke=0)

    c.setFont("Arimo-Bold", 23)
    c.setFillColor(HexColor(f"#{NAVY}"))
    c.drawString(0.48 * inch, height - 0.72 * inch, "Proceso Unificado aplicado al rediseño de Frogger")
    c.setFont("Arimo", 10)
    c.setFillColor(HexColor(f"#{MID}"))
    c.drawString(
        0.5 * inch,
        height - 1.0 * inch,
        "Flujo iterativo para modernizar el juego retro y llevarlo a PC y dispositivos móviles",
    )
    rounded(c, 10.1 * inch, height - 0.9 * inch, 2.72 * inch, 0.58 * inch, NAVY, NAVY)
    c.setFillColor(white)
    c.setFont("Arimo-Bold", 8.8)
    c.drawCentredString(11.46 * inch, height - 0.56 * inch, "CRISTOPHER ALEXANDER")
    c.drawCentredString(11.46 * inch, height - 0.75 * inch, "BAÑARES SUAREZ")

    card_x = [0.42, 3.65, 6.88, 10.11]
    card_y = 1.41 * inch
    card_h = 4.62 * inch
    for idx, (phase, x_in) in enumerate(zip(phases, card_x)):
        x = x_in * inch
        rounded(c, x, card_y, 2.80 * inch, card_h, WHITE, LINE)
        rounded(c, x, card_y + card_h - 0.72 * inch, 2.80 * inch, 0.72 * inch, phase["color"], phase["color"])
        c.setFillColor(white)
        c.setFont("Arimo-Bold", 12)
        c.drawString(x + 0.18 * inch, card_y + card_h - 0.44 * inch, phase["number"])
        c.setFont("Arimo-Bold", 13)
        c.drawString(x + 0.72 * inch, card_y + card_h - 0.44 * inch, phase["name"])
        y = card_y + card_h - 0.98 * inch
        y = draw_wrapped(
            c,
            phase["objective"],
            x + 0.17 * inch,
            y,
            2.45 * inch,
            "Arimo-Bold",
            9.3,
            NAVY,
            11,
            3,
        )
        c.setFont("Arimo-Bold", 7.5)
        c.setFillColor(HexColor(f"#{phase['color']}"))
        c.drawString(x + 0.17 * inch, y - 5, "ACTIVIDADES")
        y -= 20
        for activity in phase["activities"]:
            c.setFont("Arimo-Bold", 8)
            c.setFillColor(HexColor(f"#{phase['color']}"))
            c.drawString(x + 0.17 * inch, y, "•")
            y = draw_wrapped(
                c,
                activity,
                x + 0.3 * inch,
                y,
                2.29 * inch,
                "Arimo",
                8.2,
                DARK,
                10,
                3,
            )
            y -= 4
        artifacts_y = card_y + 0.7 * inch
        c.setFont("Arimo-Bold", 7.5)
        c.setFillColor(HexColor(f"#{phase['color']}"))
        c.drawString(x + 0.17 * inch, artifacts_y, "ARTEFACTOS")
        for tag_idx, artifact in enumerate(phase["artifacts"]):
            py = artifacts_y - (0.25 + tag_idx * 0.23) * inch
            rounded(c, x + 0.15 * inch, py, 2.46 * inch, 0.19 * inch, phase["pale"], phase["pale"], 4)
            c.setFont("Arimo-Bold", 6.8)
            c.setFillColor(HexColor(f"#{phase['color']}"))
            c.drawString(x + 0.24 * inch, py + 0.055 * inch, artifact)

        if idx < 3:
            start = (x_in + 2.82) * inch
            end = (card_x[idx + 1] - 0.05) * inch
            arrow_y = height - 1.82 * inch
            c.setStrokeColor(HexColor(f"#{MID}"))
            c.setLineWidth(1.7)
            c.line(start, arrow_y, end, arrow_y)
            c.setFillColor(HexColor(f"#{MID}"))
            c.line(end - 5, arrow_y + 3, end, arrow_y)
            c.line(end - 5, arrow_y - 3, end, arrow_y)

    for label, x in [("Alcance aprobado", 3.07), ("Arquitectura base", 6.29), ("Incrementos listos", 9.48)]:
        c.setFont("Arimo-Bold", 6.5)
        c.setFillColor(HexColor(f"#{MID}"))
        c.drawCentredString((x + 0.67) * inch, height - 1.29 * inch, label)

    rounded(c, 0.42 * inch, 0.38 * inch, 12.49 * inch, 0.84 * inch, NAVY, NAVY)
    c.setFillColor(white)
    c.setFont("Arimo-Bold", 9)
    c.drawString(0.67 * inch, 0.86 * inch, "ITERACIONES")
    c.setFont("Arimo-Bold", 9.2)
    c.drawCentredString(
        4.6 * inch,
        0.83 * inch,
        "Requisitos  →  diseño  →  implementación  →  pruebas  →  incremento jugable",
    )
    c.setFont("Arimo-Bold", 8)
    c.setFillColor(HexColor("#CFE8D1"))
    c.drawCentredString(
        10.12 * inch,
        0.82 * inch,
        "↶  Beta y pruebas PC/móvil devuelven ajustes a Construcción y Elaboración",
    )
    c.setFont("Arimo", 6.7)
    c.setFillColor(HexColor(f"#{MID}"))
    c.drawCentredString(
        width / 2,
        0.17 * inch,
        "Diseño de Software · Semana 2  |  El PU valida temprano y permite adaptar requisitos; no espera al final como un proceso en cascada.",
    )
    c.save()


if __name__ == "__main__":
    build_pptx()
    build_pdf()
    print(PPTX_OUTPUT)
    print(PDF_OUTPUT)
