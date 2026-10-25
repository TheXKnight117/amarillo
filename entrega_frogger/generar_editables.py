"""Genera las versiones editables del informe y del diagrama de Frogger Mobile."""

from pathlib import Path
from xml.etree.ElementTree import Element, SubElement, ElementTree

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
DOCX_PATH = BASE / "Informe_Frogger_Mobile_Editable.docx"
DRAWIO_PATH = BASE / "Diagrama_Clases_Frogger_Mobile.drawio"
DIAGRAM_PATH = BASE / "Diagrama_Clases_Frogger_Mobile.png"
AUTHOR = "Cristopher Alexander Bañares Suarez"


def set_cell_shading(cell, fill):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_footer(section, page_label=""):
    paragraph = section.footer.paragraphs[0]
    paragraph.text = f"{AUTHOR}  |  Diseño de Software{page_label}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(130, 154, 177)


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(23, 107, 135) if level > 1 else RGBColor(16, 42, 67)
    return paragraph


def add_bullets(document, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        paragraph = document.add_paragraph(style=style)
        paragraph.add_run(item)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, heading in enumerate(headers):
        header_cells[index].text = heading
        set_cell_shading(header_cells[index], "176B87")
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in header_cells[index].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cells[col_index].text = value
            cells[col_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if row_index % 2:
                set_cell_shading(cells[col_index], "F5F8FB")
            for paragraph in cells[col_index].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        if widths:
            for col_index, width in enumerate(widths):
                cells[col_index].width = Cm(width)
    document.add_paragraph()
    return table


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    add_footer(section)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        styles[name].font.name = "Arial"

    for _ in range(5):
        document.add_paragraph()
    label = document.add_paragraph("ACTIVIDAD FORMATIVA · SEMANA 4")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.runs[0].font.bold = True
    label.runs[0].font.color.rgb = RGBColor(23, 107, 135)
    label.runs[0].font.size = Pt(12)
    title = document.add_paragraph("Diseño de un diagrama de clases UML para Frogger Mobile")
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(16, 42, 67)
    subtitle = document.add_paragraph("Informe de modelado orientado a objetos")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(72, 101, 129)
    document.add_paragraph()
    identity = add_table(
        document,
        ["Datos de la entrega", "Información"],
        [
            ("Estudiante", AUTHOR),
            ("Asignatura", "Diseño de Software"),
            ("Producto", "Informe y diagrama de clases UML"),
            ("Modalidad", "Trabajo individual"),
        ],
        [5, 11],
    )
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    add_heading(document, "1. Introducción y alcance")
    document.add_paragraph(
        "El presente informe transforma los elementos relevantes del videojuego Frogger Mobile en un "
        "modelo estructural de clases UML. El propósito es representar las responsabilidades, los datos "
        "y las colaboraciones necesarias para una versión móvil del juego. En ella, una rana cruza un "
        "escenario dividido en carriles, evita vehículos, utiliza troncos como apoyo y alcanza una o más "
        "metas. El modelo describe la estructura del dominio sin incorporar detalles de interfaz gráfica, "
        "almacenamiento o plataforma que no son necesarios para explicar la mecánica principal."
    )
    add_heading(document, "1.1 Objetivo", 2)
    document.add_paragraph(
        "Diseñar un diagrama de clases con atributos, métodos, visibilidad, multiplicidades y relaciones "
        "UML, aplicando abstracción, encapsulamiento, modularidad y jerarquía simple."
    )
    add_heading(document, "1.2 Supuestos de diseño", 2)
    add_bullets(
        document,
        [
            "Una instancia de Juego controla una rana activa, un marcador y uno o más niveles.",
            "Cada nivel se compone de uno o más carriles y una o más metas.",
            "Los carriles agrupan cero o más obstáculos; estos pueden ser vehículos o troncos.",
            "La rana pierde una vida al colisionar con un vehículo o al no superar una zona de riesgo.",
            "Un tronco puede sostener y transportar temporalmente a la rana.",
            "EstadoJuego, Direccion, TipoCarril y TipoVehiculo restringen valores válidos, pero no se desarrollan como clases para evitar detalle innecesario.",
        ],
    )
    add_heading(document, "1.3 Proceso aplicado", 2)
    add_bullets(
        document,
        [
            "Identificación de sustantivos y responsabilidades del caso.",
            "Selección de clases y eliminación de elementos redundantes.",
            "Asignación de atributos y operaciones coherentes.",
            "Definición de herencia, composición, agregación, asociación y dependencia.",
            "Revisión de visibilidad, multiplicidades, legibilidad y modularidad.",
        ],
        numbered=True,
    )
    document.add_page_break()

    add_heading(document, "2. Identificación de clases")
    class_rows = [
        ("Juego", "Coordinar el ciclo general y el estado de la partida.", "estado, puntaje, vidasRestantes", "iniciar, pausar, reanudar, finalizar, actualizar"),
        ("Marcador", "Administrar el puntaje actual y máximo.", "puntajeActual, puntajeMaximo", "sumarPuntos, reiniciar, obtenerPuntaje"),
        ("Nivel", "Representar una etapa, su dificultad y condición de término.", "numero, dificultad, tiempoLimite", "iniciar, aumentarDificultad, comprobarFinalizacion"),
        ("Carril", "Organizar una franja del escenario y sus obstáculos.", "indice, tipo, direccion", "agregarObstaculo, actualizar"),
        ("Meta", "Representar un destino disponible u ocupado.", "identificador, ocupada", "ocupar, liberar, estaDisponible"),
        ("EntidadMovil", "Abstraer estado y operaciones compartidas por actores que se desplazan.", "posicionX, posicionY, velocidad, activo", "mover, actualizar, obtenerPosicion"),
        ("Rana", "Representar al personaje controlado y sus acciones.", "vidas, posicionInicial", "saltar, perderVida, reiniciarPosicion, llegoAMeta"),
        ("Obstaculo", "Definir el comportamiento común de objetos de riesgo o apoyo.", "direccion", "colisionar, salirDePantalla"),
        ("Vehiculo", "Especializar un obstáculo dañino terrestre.", "tipo, danio", "atropellar"),
        ("Tronco", "Especializar un obstáculo acuático que puede transportar a la rana.", "longitud", "transportar, puedeSostener"),
    ]
    add_table(document, ["Clase", "Responsabilidad", "Atributos", "Métodos"], class_rows, [2.5, 6.2, 4.2, 4.5])
    add_heading(document, "2.1 Coherencia de atributos y operaciones", 2)
    document.add_paragraph(
        "Los atributos describen el estado que cada clase necesita conservar. Las operaciones expresan "
        "comportamientos propios y evitan que otras entidades manipulen directamente sus datos. Por "
        "ejemplo, el puntaje se modifica mediante sumarPuntos() y la disponibilidad de una meta mediante "
        "ocupar() o liberar(). EntidadMovil y Obstaculo son abstractas: organizan características comunes, "
        "mientras el juego utiliza instancias concretas como Rana, Vehiculo y Tronco."
    )
    document.add_page_break()

    add_heading(document, "3. Aplicación de principios orientados a objetos")
    principles = [
        ("Abstracción", "Se conservan solo los elementos esenciales de la mecánica. Rana conoce sus vidas y posición inicial, pero no incluye aspectos visuales que corresponden a otros subsistemas."),
        ("Encapsulamiento", "Los atributos privados (−) impiden cambios externos sin control. Los métodos públicos (+) forman la interfaz y los datos reutilizables por subclases son protegidos (#)."),
        ("Modularidad", "Las clases se separan en Núcleo del juego, Escenario y Actores móviles. Cada módulo concentra una responsabilidad y mantiene relaciones explícitas."),
        ("Jerarquía", "EntidadMovil generaliza a Rana y Obstaculo; Vehiculo y Tronco especializan a Obstaculo. Así se reutilizan posición, velocidad y actualización."),
    ]
    for name, text in principles:
        add_heading(document, name, 2)
        document.add_paragraph(text)
    add_heading(document, "3.1 Decisiones de visibilidad", 2)
    add_table(
        document,
        ["Símbolo", "Uso", "Justificación"],
        [
            ("− Privado", "Estado, puntaje, vidas, dificultad, ocupación y tipos.", "Protege invariantes y obliga a usar operaciones controladas."),
            ("# Protegido", "Posición, velocidad, estado activo y dirección.", "Permite que las especializaciones implementen su movimiento sin exponer datos."),
            ("+ Público", "Acciones como iniciar, saltar, colisionar, ocupar y actualizar.", "Define los servicios mediante los cuales colaboran los objetos."),
        ],
        [3, 7, 7],
    )
    document.add_page_break()

    add_heading(document, "4. Relaciones y multiplicidades UML")
    relations = [
        ("Juego–Marcador", "Composición 1 a 1", "El marcador pertenece a una partida y su ciclo de vida depende de ella."),
        ("Juego–Rana", "Composición 1 a 1", "Cada partida controla exactamente una rana activa."),
        ("Juego–Nivel", "Composición 1 a 1..*", "Un juego contiene uno o más niveles."),
        ("Nivel–Carril", "Composición 1 a 1..*", "Un nivel se construye con uno o más carriles."),
        ("Nivel–Meta", "Composición 1 a 1..*", "Las metas son partes propias del nivel."),
        ("Carril–Obstaculo", "Agregación 1 a 0..*", "Un carril agrupa obstáculos, pero puede existir temporalmente vacío."),
        ("EntidadMovil–Rana/Obstaculo", "Generalización", "Las subclases heredan posición, velocidad y actualización."),
        ("Obstaculo–Vehiculo/Tronco", "Generalización", "Cada tipo implementa un efecto particular sobre la rana."),
        ("Rana–Tronco", "Asociación 1 a 0..1", "La rana puede estar apoyada en ningún tronco o en uno."),
        ("Rana–Meta", "Dependencia", "La rana utiliza temporalmente una meta para completar el recorrido."),
    ]
    add_table(document, ["Relación", "Notación", "Interpretación"], relations, [4.5, 4, 8])
    add_heading(document, "4.1 Ejemplo de colaboración", 2)
    add_bullets(
        document,
        [
            "Juego.iniciar() activa el nivel y reinicia el marcador.",
            "La entrada del usuario provoca Rana.saltar(direccion).",
            "Juego.actualizar(deltaTiempo) solicita la actualización de las entidades móviles.",
            "Un vehículo que coincide con la rana ejecuta atropellar(rana), y la rana procesa perderVida().",
            "Un tronco que sostiene a la rana ejecuta transportar(rana).",
            "Al alcanzar una meta disponible se ejecuta Meta.ocupar() y se suman puntos.",
        ],
        numbered=True,
    )
    document.add_page_break()

    add_heading(document, "5. Diagrama de clases")
    document.add_paragraph(
        "La figura presenta las clases organizadas por módulos, con sus atributos, operaciones, niveles "
        "de visibilidad y relaciones principales."
    )
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(DIAGRAM_PATH), width=Cm(17.2))
    caption = document.add_paragraph("Figura 1. Diagrama de clases UML propuesto para Frogger Mobile. Elaboración propia.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)
    document.add_paragraph(
        "Los rombos negros representan composición; el rombo blanco, agregación; los triángulos vacíos, "
        "herencia; y la línea discontinua, dependencia."
    )
    document.add_page_break()

    add_heading(document, "6. Conclusiones y verificación")
    document.add_paragraph(
        "El modelo representa la estructura central de Frogger Mobile mediante clases con responsabilidades "
        "delimitadas. La abstracción evita detalles ajenos a la mecánica; el encapsulamiento protege el "
        "estado; la modularidad organiza el sistema en componentes comprensibles; y la jerarquía reutiliza "
        "las características comunes de las entidades móviles."
    )
    document.add_paragraph(
        "Las relaciones UML muestran la pertenencia fuerte de niveles y elementos estructurales, junto con "
        "las interacciones temporales entre la rana, los troncos y las metas. Las multiplicidades eliminan "
        "ambigüedades y precisan cuántas instancias participan en cada colaboración."
    )
    add_heading(document, "6.1 Revisión según los criterios de evaluación", 2)
    add_table(
        document,
        ["Criterio", "Cumplimiento"],
        [
            ("Clases y atributos", "Diez clases clave con atributos y métodos vinculados a responsabilidades concretas."),
            ("Abstracción y encapsulamiento", "Se excluyen detalles irrelevantes y se aplican visibilidades pública, privada y protegida."),
            ("Modularidad y relaciones", "Tres paquetes, generalización, composición, agregación, asociación, dependencia y multiplicidades."),
            ("Claridad y legibilidad", "Clases agrupadas por función, leyenda y etiquetas de relaciones."),
            ("Formato solicitado", "Informe PDF y DOCX; diagrama PNG, PlantUML y draw.io."),
        ],
        [5, 12],
    )
    add_heading(document, "Referencias bibliográficas", 2)
    references = [
        "Booch, G., Rumbaugh, J. y Jacobson, I. (2006). El lenguaje unificado de modelado: guía del usuario. Pearson.",
        "Pressman, R. (2010). Ingeniería de software: un enfoque práctico. McGraw-Hill.",
        "Schmuller, J. (2004). Aprenda UML en 24 horas. Pearson Educación.",
        "Santo Tomás Online. (s. f.). Diseño de software: Apunte semana 4 [Material de asignatura].",
        "Sommerville, I. (2011). Ingeniería del software (9.ª ed.). Pearson.",
    ]
    for reference in references:
        document.add_paragraph(reference)

    core = document.core_properties
    core.title = "Diseño de un diagrama de clases UML para Frogger Mobile"
    core.subject = "Actividad formativa, semana 4"
    core.author = AUTHOR
    document.save(DOCX_PATH)


def add_vertex(root, identifier, value, x, y, width, height, style):
    cell = SubElement(root, "mxCell", id=identifier, value=value, style=style, vertex="1", parent="1")
    SubElement(cell, "mxGeometry", x=str(x), y=str(y), width=str(width), height=str(height), **{"as": "geometry"})
    return cell


def add_edge(root, identifier, source, target, value="", style=""):
    base_style = "edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;"
    cell = SubElement(
        root,
        "mxCell",
        id=identifier,
        value=value,
        style=base_style + style,
        edge="1",
        parent="1",
        source=source,
        target=target,
    )
    SubElement(cell, "mxGeometry", relative="1", **{"as": "geometry"})
    return cell


def uml_text(name, attributes, methods, abstract=False):
    title = f"<i>«abstract» {name}</i>" if abstract else f"<b>{name}</b>"
    return (
        f"<div style='text-align:center;font-size:14px'>{title}</div>"
        f"<hr><div style='text-align:left'>{'<br>'.join(attributes)}</div>"
        f"<hr><div style='text-align:left'>{'<br>'.join(methods)}</div>"
    )


def build_drawio():
    mxfile = Element(
        "mxfile",
        host="app.diagrams.net",
        modified="2026-10-25T16:00:00.000Z",
        agent="Cursor",
        version="29.0.0",
        type="device",
    )
    diagram = SubElement(mxfile, "diagram", id="frogger-mobile", name="Frogger Mobile")
    model = SubElement(
        diagram,
        "mxGraphModel",
        dx="1800",
        dy="1200",
        grid="1",
        gridSize="10",
        guides="1",
        tooltips="1",
        connect="1",
        arrows="1",
        fold="1",
        page="1",
        pageScale="1",
        pageWidth="1800",
        pageHeight="1200",
        math="0",
        shadow="0",
    )
    root = SubElement(model, "root")
    SubElement(root, "mxCell", id="0")
    SubElement(root, "mxCell", id="1", parent="0")

    package_style = "rounded=1;whiteSpace=wrap;html=1;dashed=1;verticalAlign=top;align=left;spacingTop=8;spacingLeft=10;fillColor=#f5f8fb;strokeColor=#9fb3c8;fontStyle=1;"
    add_vertex(root, "pkg-core", "Núcleo del juego", 20, 30, 390, 610, package_style)
    add_vertex(root, "pkg-scene", "Escenario", 430, 30, 390, 940, package_style)
    add_vertex(root, "pkg-actors", "Actores móviles", 840, 30, 900, 940, package_style)

    class_style = "rounded=0;whiteSpace=wrap;html=1;align=left;verticalAlign=top;spacing=8;fillColor=#ffffff;strokeColor=#243b53;fontColor=#102a43;"
    abstract_style = class_style + "fillColor=#eef2f6;"
    classes = {
        "juego": ("Juego", ["− estado: EstadoJuego", "− puntaje: int", "− vidasRestantes: int"], ["+ iniciar(): void", "+ pausar(): void", "+ reanudar(): void", "+ finalizar(): void", "+ actualizar(dt: float): void"], 50, 80, 320, 245, False),
        "marcador": ("Marcador", ["− puntajeActual: int", "− puntajeMaximo: int"], ["+ sumarPuntos(cantidad: int): void", "+ reiniciar(): void", "+ obtenerPuntaje(): int"], 50, 390, 320, 190, False),
        "nivel": ("Nivel", ["− numero: int", "− dificultad: int", "− tiempoLimite: int"], ["+ iniciar(): void", "+ aumentarDificultad(): void", "+ comprobarFinalizacion(): boolean"], 465, 80, 320, 215, False),
        "carril": ("Carril", ["− indice: int", "− tipo: TipoCarril", "− direccion: Direccion"], ["+ agregarObstaculo(o: Obstaculo): void", "+ actualizar(): void"], 465, 370, 320, 195, False),
        "meta": ("Meta", ["− identificador: int", "− ocupada: boolean"], ["+ ocupar(): void", "+ liberar(): void", "+ estaDisponible(): boolean"], 465, 660, 320, 190, False),
        "entidad": ("EntidadMovil", ["# posicionX: float", "# posicionY: float", "# velocidad: float", "# activo: boolean"], ["+ mover(): void", "+ actualizar(dt: float): void", "+ obtenerPosicion(): Punto"], 1010, 70, 350, 235, True),
        "rana": ("Rana", ["− vidas: int", "− posicionInicial: Punto"], ["+ saltar(d: Direccion): void", "+ perderVida(): void", "+ reiniciarPosicion(): void", "+ llegoAMeta(): boolean"], 875, 390, 345, 220, False),
        "obstaculo": ("Obstaculo", ["# direccion: Direccion"], ["+ colisionar(rana: Rana): void", "+ salirDePantalla(): boolean"], 1380, 380, 330, 175, True),
        "vehiculo": ("Vehiculo", ["− tipo: TipoVehiculo", "− danio: int"], ["+ atropellar(rana: Rana): void"], 1030, 720, 300, 155, False),
        "tronco": ("Tronco", ["− longitud: float"], ["+ transportar(rana: Rana): void", "+ puedeSostener(rana: Rana): boolean"], 1400, 720, 310, 165, False),
    }
    for identifier, (name, attrs, methods, x, y, width, height, abstract) in classes.items():
        add_vertex(root, identifier, uml_text(name, attrs, methods, abstract), x, y, width, height, abstract_style if abstract else class_style)

    composition = "startArrow=diamond;startFill=1;endArrow=none;"
    aggregation = "startArrow=diamond;startFill=0;endArrow=none;"
    inheritance = "endArrow=block;endFill=0;"
    dependency = "dashed=1;endArrow=open;endFill=0;"
    association = "endArrow=none;"
    add_edge(root, "e1", "juego", "marcador", "1      1", composition)
    add_edge(root, "e2", "juego", "nivel", "1      1..*", composition)
    add_edge(root, "e3", "juego", "rana", "1 controla 1", composition)
    add_edge(root, "e4", "nivel", "carril", "1      1..*", composition)
    add_edge(root, "e5", "nivel", "meta", "1      1..*", composition)
    add_edge(root, "e6", "carril", "obstaculo", "1 agrupa 0..*", aggregation)
    add_edge(root, "e7", "rana", "entidad", "", inheritance)
    add_edge(root, "e8", "obstaculo", "entidad", "", inheritance)
    add_edge(root, "e9", "vehiculo", "obstaculo", "", inheritance)
    add_edge(root, "e10", "tronco", "obstaculo", "", inheritance)
    add_edge(root, "e11", "rana", "tronco", "0..1 se apoya en", association)
    add_edge(root, "e12", "rana", "meta", "intenta alcanzar", dependency)

    ElementTree(mxfile).write(DRAWIO_PATH, encoding="utf-8", xml_declaration=True)


if __name__ == "__main__":
    build_docx()
    build_drawio()
    print(DOCX_PATH)
    print(DRAWIO_PATH)
