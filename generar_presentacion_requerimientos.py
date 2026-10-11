from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepInFrame,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


PPTX_OUTPUT = Path("S2_Bañares_Requerimientos_Registro_Usuario.pptx")
PDF_OUTPUT = Path("S2_Bañares_Requerimientos_Registro_Usuario.pdf")

NAVY = "102A43"
BLUE = "1677FF"
CYAN = "1CB5E0"
TEAL = "0EAD8E"
ORANGE = "FF9F43"
PALE_BLUE = "EAF4FF"
PALE_TEAL = "E7F8F4"
PALE_ORANGE = "FFF3E5"
LIGHT = "F5F8FC"
MID = "627D98"
DARK = "243B53"
WHITE = "FFFFFF"
RED = "D64545"


rf = [
    (
        "RF-01",
        "Capturar datos",
        "El sistema deberá solicitar nombre, correo electrónico, contraseña y aceptación de términos.",
        "Define los datos y la acción que el sistema debe ejecutar.",
    ),
    (
        "RF-02",
        "Validar información",
        "El sistema deberá comprobar campos obligatorios, formato del correo, correo único y política de contraseña.",
        "Es funcional porque describe validaciones concretas del sistema.",
    ),
    (
        "RF-03",
        "Crear la cuenta",
        "Si los datos son válidos, el sistema deberá registrar la cuenta y asociarla al correo ingresado.",
        "Representa el servicio principal solicitado por el usuario.",
    ),
    (
        "RF-04",
        "Verificar el correo",
        "El sistema deberá enviar un enlace de verificación y activar la cuenta cuando el usuario lo utilice.",
        "Indica un comportamiento y un resultado observable.",
    ),
]

rnf = [
    (
        "RNF-01",
        "Usabilidad",
        "Un usuario nuevo deberá completar el formulario en un máximo de 2 minutos; los errores deberán indicar cómo corregirse.",
        "Mide facilidad de uso y comprensión.",
    ),
    (
        "RNF-02",
        "Seguridad",
        "Los datos deberán viajar mediante TLS 1.2 o superior y las contraseñas se almacenarán con hash seguro y salt.",
        "Restringe cómo se protegen los datos.",
    ),
    (
        "RNF-03",
        "Rendimiento",
        "El 95 % de las solicitudes de registro deberá responder en 2 segundos o menos con 100 usuarios concurrentes.",
        "Establece una medida verificable de velocidad.",
    ),
    (
        "RNF-04",
        "Confiabilidad",
        "El servicio de registro deberá alcanzar 99,9 % de disponibilidad mensual y evitar cuentas duplicadas ante reintentos.",
        "Define continuidad y comportamiento ante fallos.",
    ),
]

matrix_rf = [
    ["RF-01", "RF", "Ingresar datos", "Formulario de registro", "CU-01 / CP-01: registro válido"],
    ["RF-02", "RF", "Validar campos", "Validación del formulario", "CU-02 / CP-02 y CP-03: datos inválidos"],
    ["RF-03", "RF", "Crear cuenta", "Persistencia de usuario", "CU-01 / CP-01: cuenta almacenada"],
    ["RF-04", "RF", "Verificar correo", "Activación de cuenta", "CU-03 / CP-04: enlace de activación"],
]

matrix_rnf = [
    ["RNF-01", "RNF", "Registro ≤ 2 min", "Experiencia de registro", "CP-05: prueba con usuario nuevo"],
    ["RNF-02", "RNF", "TLS + hash seguro", "Protección de datos", "CP-06: revisar tráfico y almacenamiento"],
    ["RNF-03", "RNF", "Respuesta ≤ 2 s", "Procesamiento de registro", "CP-07: carga con 100 usuarios"],
    ["RNF-04", "RNF", "99,9 % y sin duplicados", "Disponibilidad e integridad", "CP-08: reintento tras interrupción"],
]


def rgb(value):
    return RGBColor.from_string(value)


def add_rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def set_text(
    shape,
    text,
    size=18,
    color=DARK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Aptos",
    margin=0.08,
    valign=MSO_ANCHOR.MIDDLE,
):
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)


def add_text(slide, x, y, w, h, text, **kwargs):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box, text, **kwargs)
    return box


def add_rich_text(slide, x, y, w, h, lines, fill=None, border=None, margin=0.18):
    if fill:
        shape = add_rect(slide, x, y, w, h, fill, True, border or fill)
    else:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(lines):
        text, size, color, bold = line[:4]
        space = line[4] if len(line) > 4 else 5
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def add_header(slide, number, title, kicker):
    add_rect(slide, 0, 0, 13.333, 0.16, BLUE)
    add_text(slide, 0.65, 0.35, 10.9, 0.25, kicker.upper(), size=9, color=BLUE, bold=True)
    add_text(slide, 0.65, 0.67, 11.8, 0.55, title, size=25, color=NAVY, bold=True)
    add_text(slide, 12.25, 0.48, 0.48, 0.48, f"{number:02}", size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    circle = slide.shapes[-1]
    circle_shape = add_rect(slide, 12.2, 0.42, 0.55, 0.55, NAVY, True)
    slide.shapes._spTree.remove(circle._element)
    slide.shapes._spTree.insert(len(slide.shapes._spTree) - 1, circle._element)
    set_text(circle_shape, f"{number:02}", size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_footer(slide):
    add_text(
        slide,
        0.65,
        7.12,
        12,
        0.16,
        "IGRV0129  |  Especificación de requerimientos  |  Semana 2",
        size=7.5,
        color=MID,
        align=PP_ALIGN.RIGHT,
    )


def add_card(slide, x, y, w, h, tag, title, body, note, accent):
    add_rect(slide, x, y, w, h, WHITE, True, "D9E2EC")
    add_rect(slide, x, y, 0.08, h, accent, True, accent)
    add_text(slide, x + 0.24, y + 0.18, 0.72, 0.28, tag, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    tag_shape = add_rect(slide, x + 0.2, y + 0.14, 0.82, 0.34, accent, True)
    set_text(tag_shape, tag, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 1.12, y + 0.13, w - 1.3, 0.38, title, size=14, color=NAVY, bold=True)
    add_text(slide, x + 0.23, y + 0.62, w - 0.46, 0.74, body, size=10.2, color=DARK, valign=MSO_ANCHOR.TOP)
    add_text(slide, x + 0.23, y + h - 0.47, w - 0.46, 0.3, note, size=8.5, color=accent, bold=True)


def add_table(slide, x, y, w, h, headers, rows, widths, font_size=9):
    table = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY)
        set_text(cell, header, size=font_size, color=WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0.05)
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(WHITE if r % 2 else LIGHT)
            set_text(
                cell,
                value,
                size=font_size - 0.4,
                color=DARK,
                bold=c == 0,
                align=PP_ALIGN.CENTER if c < 2 else PP_ALIGN.LEFT,
                margin=0.06,
                valign=MSO_ANCHOR.MIDDLE,
            )
    return table


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Portada
    slide = prs.slides.add_slide(blank)
    add_rect(slide, 0, 0, 13.333, 7.5, NAVY)
    add_rect(slide, 0, 0, 0.2, 7.5, CYAN)
    add_text(slide, 0.85, 0.78, 3.1, 0.35, "ACTIVIDAD FORMATIVA · SEMANA 2", size=11, color=CYAN, bold=True)
    add_text(
        slide,
        0.85,
        1.38,
        9.9,
        1.65,
        "Requerimientos para el\nregistro de usuarios",
        size=34,
        color=WHITE,
        bold=True,
        valign=MSO_ANCHOR.TOP,
    )
    add_text(
        slide,
        0.88,
        3.24,
        8.8,
        0.65,
        "Plataforma de compras en línea",
        size=19,
        color="B8D9FF",
    )
    add_rect(slide, 0.85, 4.38, 5.7, 1.65, "173A5E", True, "2B557B")
    add_text(slide, 1.1, 4.63, 5.2, 0.3, "CRISTOPHER ALEXANDER BAÑARES SUAREZ", size=11, color=WHITE, bold=True)
    add_text(slide, 1.1, 5.03, 5.2, 0.28, "Carrera: ______________________________", size=10.5, color="D9EAFB")
    add_text(slide, 1.1, 5.39, 5.2, 0.28, "Asignatura: Especificación de requerimientos (IGRV0129)", size=10.5, color="D9EAFB")
    # Ilustración simple del flujo
    for i, (label, accent) in enumerate([("Datos", BLUE), ("Validación", CYAN), ("Cuenta", TEAL)]):
        x = 8.1 + i * 1.45
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(4.5), Inches(1.02), Inches(1.02))
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(accent)
        circle.line.color.rgb = rgb(accent)
        set_text(circle, str(i + 1), size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, x - 0.15, 5.65, 1.32, 0.3, label, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        if i < 2:
            add_text(slide, x + 1.02, 4.79, 0.4, 0.3, "→", size=19, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 2. Introducción
    slide = prs.slides.add_slide(blank)
    add_header(slide, 2, "¿Qué se analizará?", "Introducción a la funcionalidad")
    add_rich_text(
        slide,
        0.7,
        1.48,
        7.1,
        4.7,
        [
            ("Registro de usuario", 19, NAVY, True, 12),
            (
                "Es el proceso mediante el cual una persona entrega sus datos, crea credenciales y confirma su identidad para obtener una cuenta en una tienda en línea.",
                13,
                DARK,
                False,
                15,
            ),
            ("¿Por qué es importante?", 15, BLUE, True, 8),
            (
                "Permite identificar al cliente, proteger su información y habilitar funciones como comprar, consultar pedidos, guardar direcciones y recuperar la contraseña.",
                12,
                DARK,
                False,
                12,
            ),
            (
                "El análisis separa lo que el sistema debe hacer (RF) de las condiciones de calidad con que debe hacerlo (RNF).",
                11,
                MID,
                True,
                0,
            ),
        ],
        fill=WHITE,
        border="D9E2EC",
        margin=0.32,
    )
    add_rect(slide, 8.35, 1.55, 4.2, 4.55, PALE_BLUE, True, "C6DDF5")
    flow = [("1", "Completar formulario"), ("2", "Validar información"), ("3", "Crear cuenta"), ("4", "Verificar correo")]
    for i, (n, label) in enumerate(flow):
        y = 1.92 + i * 0.92
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.75), Inches(y), Inches(0.48), Inches(0.48))
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(BLUE if i < 3 else TEAL)
        circle.line.color.rgb = rgb(BLUE if i < 3 else TEAL)
        set_text(circle, n, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, 9.48, y - 0.02, 2.55, 0.5, label, size=12, color=NAVY, bold=i == 3)
        if i < 3:
            add_text(slide, 8.82, y + 0.48, 0.32, 0.35, "↓", size=14, color=MID, align=PP_ALIGN.CENTER)
    add_footer(slide)

    # 3. RF
    slide = prs.slides.add_slide(blank)
    add_header(slide, 3, "Requerimientos funcionales", "Qué debe hacer el sistema")
    coords = [(0.7, 1.45), (6.85, 1.45), (0.7, 4.15), (6.85, 4.15)]
    for item, (x, y) in zip(rf, coords):
        add_card(slide, x, y, 5.78, 2.25, *item, BLUE)
    add_footer(slide)

    # 4. RNF
    slide = prs.slides.add_slide(blank)
    add_header(slide, 4, "Requerimientos no funcionales", "Cómo debe operar el sistema")
    coords = [(0.7, 1.45), (6.85, 1.45), (0.7, 4.15), (6.85, 4.15)]
    accents = [TEAL, ORANGE, BLUE, "7B61FF"]
    for item, (x, y), accent in zip(rnf, coords, accents):
        add_card(slide, x, y, 5.78, 2.25, *item, accent)
    add_footer(slide)

    # 5. Relación
    slide = prs.slides.add_slide(blank)
    add_header(slide, 5, "Relación entre requisitos y funcionalidad", "Operatividad + calidad")
    add_rich_text(
        slide,
        0.75,
        1.58,
        4.05,
        4.65,
        [
            ("REQUERIMIENTOS FUNCIONALES", 12, BLUE, True, 14),
            ("RF-01 captura los datos.", 12, DARK, False, 9),
            ("RF-02 evita información inválida.", 12, DARK, False, 9),
            ("RF-03 crea la cuenta.", 12, DARK, False, 9),
            ("RF-04 confirma la identidad.", 12, DARK, False, 16),
            ("Aseguran la operatividad del proceso.", 11, BLUE, True, 0),
        ],
        fill=PALE_BLUE,
        border="C6DDF5",
        margin=0.3,
    )
    add_rich_text(
        slide,
        8.52,
        1.58,
        4.05,
        4.65,
        [
            ("REQUERIMIENTOS NO FUNCIONALES", 12, TEAL, True, 14),
            ("RNF-01 facilita completar el registro.", 12, DARK, False, 9),
            ("RNF-02 protege los datos.", 12, DARK, False, 9),
            ("RNF-03 evita esperas excesivas.", 12, DARK, False, 9),
            ("RNF-04 mantiene el servicio confiable.", 12, DARK, False, 16),
            ("Aseguran la calidad del proceso.", 11, TEAL, True, 0),
        ],
        fill=PALE_TEAL,
        border="B8E6DB",
        margin=0.3,
    )
    add_text(slide, 4.72, 3.05, 0.6, 0.5, "→", size=25, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 8.0, 3.05, 0.6, 0.5, "←", size=25, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    result = add_rect(slide, 5.25, 2.28, 2.8, 2.08, NAVY, True, NAVY)
    set_text(
        result,
        "REGISTRO\nCOMPLETO,\nSEGURO Y ÁGIL",
        size=16,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        4.55,
        4.72,
        4.22,
        0.8,
        "Los RF permiten que exista la función;\nlos RNF determinan si funciona con calidad.",
        size=11,
        color=MID,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    # 6 y 7. Matriz
    headers = ["ID", "Tipo", "Descripción breve", "Funcionalidad asociada", "Caso de uso / prueba sugerida"]
    widths = [0.9, 0.72, 2.25, 2.65, 4.62]
    slide = prs.slides.add_slide(blank)
    add_header(slide, 6, "Matriz de trazabilidad · RF", "Cobertura funcional")
    add_text(
        slide,
        0.72,
        1.28,
        11.85,
        0.45,
        "Cada requerimiento queda conectado con una parte del registro y con una forma concreta de comprobarlo.",
        size=11,
        color=MID,
    )
    add_table(slide, 0.72, 1.88, 11.14, 3.92, headers, matrix_rf, widths, 9)
    add_rich_text(
        slide,
        0.72,
        6.05,
        11.85,
        0.72,
        [
            (
                "Casos de uso: CU-01 Registrar usuario · CU-02 Corregir datos inválidos · CU-03 Verificar cuenta",
                9.5,
                NAVY,
                True,
                0,
            )
        ],
        fill=PALE_BLUE,
        border="C6DDF5",
        margin=0.18,
    )
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    add_header(slide, 7, "Matriz de trazabilidad · RNF", "Cobertura de calidad")
    add_text(
        slide,
        0.72,
        1.28,
        11.85,
        0.45,
        "Las condiciones de calidad también se vinculan con pruebas medibles, no solo con descripciones generales.",
        size=11,
        color=MID,
    )
    add_table(slide, 0.72, 1.88, 11.14, 3.92, headers, matrix_rnf, widths, 9)
    add_rich_text(
        slide,
        0.72,
        6.05,
        11.85,
        0.72,
        [
            (
                "Trazabilidad: requisito → funcionalidad afectada → caso de prueba → evidencia de cumplimiento",
                9.5,
                TEAL,
                True,
                0,
            )
        ],
        fill=PALE_TEAL,
        border="B8E6DB",
        margin=0.18,
    )
    add_footer(slide)

    # 8. Pruebas
    slide = prs.slides.add_slide(blank)
    add_header(slide, 8, "Casos de prueba principales", "Cómo se comprobará")
    test_rows = [
        ["CP-01", "Datos válidos", "La cuenta se crea y se envía correo.", "RF-01, RF-03, RF-04"],
        ["CP-02", "Correo duplicado", "Se rechaza y se informa el problema.", "RF-02"],
        ["CP-03", "Contraseña débil", "Se indican las reglas incumplidas.", "RF-02"],
        ["CP-04", "Enlace de verificación", "La cuenta queda activa una sola vez.", "RF-04"],
        ["CP-05–08", "Calidad", "Se cumplen tiempo, seguridad, carga y fiabilidad.", "RNF-01 a RNF-04"],
    ]
    add_table(
        slide,
        0.72,
        1.58,
        11.8,
        4.65,
        ["Prueba", "Escenario", "Resultado esperado", "Cobertura"],
        test_rows,
        [1.08, 2.25, 5.55, 2.92],
        10,
    )
    add_text(
        slide,
        0.85,
        6.42,
        11.5,
        0.4,
        "Resultado: todos los requerimientos cuentan con una verificación asociada.",
        size=11,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    # 9. Fuentes
    slide = prs.slides.add_slide(blank)
    add_header(slide, 9, "Respaldo bibliográfico", "Fuentes utilizadas")
    sources = [
        (
            "FUENTE PRINCIPAL",
            "Sommerville, I. (2011). Ingeniería del Software (9.ª ed.). Pearson Educación. Capítulo 4, pp. 101–111.",
            BLUE,
            PALE_BLUE,
        ),
        (
            "ESTÁNDAR DE REQUISITOS",
            "ISO/IEC/IEEE. (2018). ISO/IEC/IEEE 29148:2018: Systems and software engineering — Life cycle processes — Requirements engineering.",
            TEAL,
            PALE_TEAL,
        ),
        (
            "SEGURIDAD DE APLICACIONES",
            "OWASP Foundation. (2021). Application Security Verification Standard 4.0.3. https://owasp.org/www-project-application-security-verification-standard/",
            ORANGE,
            PALE_ORANGE,
        ),
        (
            "APUNTE DE LA SEMANA",
            "Santo Tomás Online. (2026). Especificación de requerimientos: Apunte semana 2.",
            "7B61FF",
            "F0ECFF",
        ),
    ]
    for i, (tag, source, accent, fill) in enumerate(sources):
        y = 1.5 + i * 1.3
        add_rect(slide, 0.75, y, 11.8, 1.02, fill, True, fill)
        add_text(slide, 1.02, y + 0.16, 2.05, 0.25, tag, size=8.5, color=accent, bold=True)
        add_text(slide, 3.05, y + 0.1, 9.05, 0.72, source, size=10.5, color=DARK, valign=MSO_ANCHOR.TOP)
    add_text(
        slide,
        0.75,
        6.85,
        11.8,
        0.22,
        "Las definiciones y la clasificación de RF/RNF se basan principalmente en Sommerville (2011, pp. 101–111).",
        size=8.5,
        color=MID,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    prs.save(PPTX_OUTPUT)


pdfmetrics.registerFont(
    TTFont("Arimo", "/usr/share/fonts/truetype/croscore/Arimo-Regular.ttf")
)
pdfmetrics.registerFont(
    TTFont("Arimo-Bold", "/usr/share/fonts/truetype/croscore/Arimo-Bold.ttf")
)


def para(text, style):
    return Paragraph(text, style)


def pdf_table(data, widths, header=True, font_size=9):
    table = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("FONTNAME", (0, 0), (-1, -1), "Arimo"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2EC")),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]
    if header:
        style += [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Arimo-Bold"),
            ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ]
        for row in range(1, len(data)):
            if row % 2 == 0:
                style.append(("BACKGROUND", (0, row), (-1, row), colors.HexColor(f"#{LIGHT}")))
    table.setStyle(TableStyle(style))
    return table


def build_pdf():
    page = landscape((13.333 * inch, 7.5 * inch))
    doc = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=page,
        leftMargin=0.62 * inch,
        rightMargin=0.62 * inch,
        topMargin=0.48 * inch,
        bottomMargin=0.42 * inch,
        title="Requerimientos para el registro de usuarios",
        author="Cristopher Alexander Bañares Suarez",
    )
    title = ParagraphStyle(
        "Title", fontName="Arimo-Bold", fontSize=25, leading=29, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=12
    )
    kicker = ParagraphStyle(
        "Kicker", fontName="Arimo-Bold", fontSize=9, leading=11, textColor=colors.HexColor(f"#{BLUE}"), spaceAfter=5
    )
    body = ParagraphStyle(
        "Body", fontName="Arimo", fontSize=11, leading=15, textColor=colors.HexColor(f"#{DARK}"), spaceAfter=9
    )
    small = ParagraphStyle(
        "Small", parent=body, fontSize=8.5, leading=11, spaceAfter=3
    )
    card_title = ParagraphStyle(
        "CardTitle", fontName="Arimo-Bold", fontSize=11, leading=14, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=6
    )
    white_title = ParagraphStyle(
        "WhiteTitle", fontName="Arimo-Bold", fontSize=29, leading=34, textColor=colors.white, alignment=TA_LEFT, spaceAfter=14
    )
    white_body = ParagraphStyle(
        "WhiteBody", fontName="Arimo", fontSize=11, leading=15, textColor=colors.HexColor("#D9EAFB"), spaceAfter=5
    )

    def page_title(k, t):
        return [para(k.upper(), kicker), para(t, title), Spacer(1, 0.08 * inch)]

    story = []
    # Cover
    cover_table = Table(
        [
            [para("ACTIVIDAD FORMATIVA · SEMANA 2", kicker)],
            [para("Requerimientos para el<br/>registro de usuarios", white_title)],
            [para("Plataforma de compras en línea", white_body)],
            [Spacer(1, 0.35 * inch)],
            [para("<b>CRISTOPHER ALEXANDER BAÑARES SUAREZ</b>", white_body)],
            [para("Carrera: ______________________________", white_body)],
            [para("Asignatura: Especificación de requerimientos (IGRV0129)", white_body)],
        ],
        colWidths=[11.9 * inch],
        rowHeights=[0.35 * inch, 1.45 * inch, 0.4 * inch, 0.55 * inch, 0.35 * inch, 0.32 * inch, 0.32 * inch],
    )
    cover_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{NAVY}")),
                ("BOX", (0, 0), (-1, -1), 0, colors.HexColor(f"#{NAVY}")),
                ("LEFTPADDING", (0, 0), (-1, -1), 28),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    story += [Spacer(1, 0.4 * inch), cover_table, PageBreak()]

    # Intro
    story += page_title("Introducción a la funcionalidad", "¿Qué se analizará?")
    intro_left = [
        para("Registro de usuario", card_title),
        para(
            "Es el proceso mediante el cual una persona entrega sus datos, crea credenciales y confirma su identidad para obtener una cuenta en una tienda en línea.",
            body,
        ),
        para("¿Por qué es importante?", card_title),
        para(
            "Permite identificar al cliente, proteger su información y habilitar funciones como comprar, consultar pedidos, guardar direcciones y recuperar la contraseña.",
            body,
        ),
        para(
            "<b>El análisis separa lo que el sistema debe hacer (RF) de las condiciones de calidad con que debe hacerlo (RNF).</b>",
            body,
        ),
    ]
    intro_right = [
        para("<b>1.</b> Completar formulario", body),
        para("<b>2.</b> Validar información", body),
        para("<b>3.</b> Crear cuenta", body),
        para("<b>4.</b> Verificar correo", body),
    ]
    intro = Table(
        [[KeepInFrame(6.7 * inch, 4.5 * inch, intro_left), KeepInFrame(4 * inch, 4.5 * inch, intro_right)]],
        colWidths=[7.2 * inch, 4.2 * inch],
        rowHeights=[4.7 * inch],
    )
    intro.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), colors.white),
                ("BACKGROUND", (1, 0), (1, 0), colors.HexColor(f"#{PALE_BLUE}")),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D9E2EC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 18),
                ("RIGHTPADDING", (0, 0), (-1, -1), 18),
                ("TOPPADDING", (0, 0), (-1, -1), 18),
            ]
        )
    )
    story += [intro, PageBreak()]

    # Requirement cards
    for kicker_text, page_name, dataset, color_hex in [
        ("Qué debe hacer el sistema", "Requerimientos funcionales", rf, BLUE),
        ("Cómo debe operar el sistema", "Requerimientos no funcionales", rnf, TEAL),
    ]:
        story += page_title(kicker_text, page_name)
        cards = []
        for ident, name, desc, why in dataset:
            card = KeepInFrame(
                5.5 * inch,
                2.0 * inch,
                [
                    para(f'<font color="#{color_hex}"><b>{ident}</b></font> · {name}', card_title),
                    para(desc, small),
                    para(f'<font color="#{color_hex}"><b>{why}</b></font>', small),
                ],
            )
            cards.append(card)
        table = Table([[cards[0], cards[1]], [cards[2], cards[3]]], colWidths=[5.75 * inch, 5.75 * inch], rowHeights=[2.25 * inch, 2.25 * inch])
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 5, colors.white),
                    ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D9E2EC")),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.white),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 14),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                    ("TOPPADDING", (0, 0), (-1, -1), 13),
                ]
            )
        )
        story += [table, PageBreak()]

    # Relationship
    story += page_title("Operatividad + calidad", "Relación entre requisitos y funcionalidad")
    left = [
        para(f'<font color="#{BLUE}"><b>REQUERIMIENTOS FUNCIONALES</b></font>', card_title),
        para("RF-01 captura los datos.<br/>RF-02 evita información inválida.<br/>RF-03 crea la cuenta.<br/>RF-04 confirma la identidad.", body),
        para(f'<font color="#{BLUE}"><b>Aseguran la operatividad del proceso.</b></font>', body),
    ]
    center = [
        para('<font color="#FFFFFF"><b>REGISTRO<br/>COMPLETO,<br/>SEGURO Y ÁGIL</b></font>', ParagraphStyle("Center", parent=body, fontSize=16, leading=21, alignment=TA_CENTER))
    ]
    right = [
        para(f'<font color="#{TEAL}"><b>REQUERIMIENTOS NO FUNCIONALES</b></font>', card_title),
        para("RNF-01 facilita completar el registro.<br/>RNF-02 protege los datos.<br/>RNF-03 evita esperas excesivas.<br/>RNF-04 mantiene el servicio confiable.", body),
        para(f'<font color="#{TEAL}"><b>Aseguran la calidad del proceso.</b></font>', body),
    ]
    rel = Table(
        [[KeepInFrame(4.2 * inch, 4.4 * inch, left), KeepInFrame(2.3 * inch, 2.2 * inch, center), KeepInFrame(4.2 * inch, 4.4 * inch, right)]],
        colWidths=[4.55 * inch, 2.5 * inch, 4.55 * inch],
        rowHeights=[4.45 * inch],
    )
    rel.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{PALE_BLUE}")),
                ("BACKGROUND", (1, 0), (1, 0), colors.HexColor(f"#{NAVY}")),
                ("BACKGROUND", (2, 0), (2, 0), colors.HexColor(f"#{PALE_TEAL}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 15),
                ("RIGHTPADDING", (0, 0), (-1, -1), 15),
                ("TOPPADDING", (0, 0), (-1, -1), 15),
                ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#D9E2EC")),
            ]
        )
    )
    story += [rel, PageBreak()]

    # Matrices
    pdf_headers = ["ID", "Tipo", "Descripción breve", "Funcionalidad asociada", "Caso de uso / prueba sugerida"]
    matrix_widths = [0.75 * inch, 0.62 * inch, 2 * inch, 2.55 * inch, 5.1 * inch]
    for n, label, rows, note in [
        ("Cobertura funcional", "Matriz de trazabilidad · RF", matrix_rf, "CU-01 Registrar usuario · CU-02 Corregir datos inválidos · CU-03 Verificar cuenta"),
        ("Cobertura de calidad", "Matriz de trazabilidad · RNF", matrix_rnf, "Requisito → funcionalidad afectada → caso de prueba → evidencia de cumplimiento"),
    ]:
        story += page_title(n, label)
        wrapped = [[para(str(v), small) for v in row] for row in [pdf_headers] + rows]
        story += [
            pdf_table(wrapped, matrix_widths, True, 8.3),
            Spacer(1, 0.3 * inch),
            para(f"<b>{note}</b>", body),
            PageBreak(),
        ]

    # Test cases
    story += page_title("Cómo se comprobará", "Casos de prueba principales")
    test_rows = [
        ["CP-01", "Datos válidos", "La cuenta se crea y se envía correo.", "RF-01, RF-03, RF-04"],
        ["CP-02", "Correo duplicado", "Se rechaza y se informa el problema.", "RF-02"],
        ["CP-03", "Contraseña débil", "Se indican las reglas incumplidas.", "RF-02"],
        ["CP-04", "Enlace de verificación", "La cuenta queda activa una sola vez.", "RF-04"],
        ["CP-05–08", "Calidad", "Se cumplen tiempo, seguridad, carga y fiabilidad.", "RNF-01 a RNF-04"],
    ]
    wrapped_tests = [[para(str(v), small) for v in row] for row in [["Prueba", "Escenario", "Resultado esperado", "Cobertura"]] + test_rows]
    story += [
        pdf_table(wrapped_tests, [1 * inch, 2.1 * inch, 5.45 * inch, 2.7 * inch], True, 9),
        Spacer(1, 0.3 * inch),
        para("<b>Resultado: todos los requerimientos cuentan con una verificación asociada.</b>", body),
        PageBreak(),
    ]

    # Bibliography
    story += page_title("Fuentes utilizadas", "Respaldo bibliográfico")
    source_rows = [
        ["FUENTE PRINCIPAL", "Sommerville, I. (2011). Ingeniería del Software (9.ª ed.). Pearson Educación. Capítulo 4, pp. 101–111."],
        ["ESTÁNDAR", "ISO/IEC/IEEE. (2018). ISO/IEC/IEEE 29148:2018: Systems and software engineering — Life cycle processes — Requirements engineering."],
        ["SEGURIDAD", "OWASP Foundation. (2021). Application Security Verification Standard 4.0.3. https://owasp.org/www-project-application-security-verification-standard/"],
        ["APUNTE", "Santo Tomás Online. (2026). Especificación de requerimientos: Apunte semana 2."],
    ]
    wrapped_sources = [[para(f"<b>{a}</b>", small), para(b, body)] for a, b in source_rows]
    source_table = Table(wrapped_sources, colWidths=[2.05 * inch, 9.35 * inch], rowHeights=[1.05 * inch] * 4)
    source_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{PALE_BLUE}")),
                ("BACKGROUND", (1, 0), (1, -1), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#D9E2EC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story += [
        source_table,
        Spacer(1, 0.22 * inch),
        para("Las definiciones y la clasificación de RF/RNF se basan principalmente en Sommerville (2011, pp. 101–111).", small),
    ]
    doc.build(story)


if __name__ == "__main__":
    build_pptx()
    build_pdf()
    print(PPTX_OUTPUT)
    print(PDF_OUTPUT)
